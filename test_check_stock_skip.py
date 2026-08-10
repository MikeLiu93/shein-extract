"""Test: check_stock.py 遇到中文新模板必须完全不写单元格，也不启动 Chrome。
新 schema 的 E/F/G 是 变体/日期/状态，老逻辑会把它们当 Status/Stock/Last Checked。
Run: python test_check_stock_skip.py
"""
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

import check_stock
import shein_scraper
from run_excel import EXPECTED_HEADERS


def _new_template_xlsx(tmp):
    wb = Workbook()
    ws = wb.active
    ws.title = "ZR1"
    for c, h in enumerate(EXPECTED_HEADERS, 1):
        ws.cell(1, c).value = h
    # 一行「已完成」——老逻辑正是拿这种行去复查库存的
    ws.append([1, "http://a", 10.0, 2.99, "Black / M",
               "2026-08-01", "Done", None, 12.99, "T", "eT", 18.19, "Black-M: 5"])
    p = Path(tmp) / "new.xlsx"
    wb.save(p)
    return p


def test_new_template_is_untouched_and_chrome_not_launched():
    tmp = tempfile.mkdtemp()
    try:
        p = _new_template_xlsx(tmp)
        before_row = [c.value for c in load_workbook(p)["ZR1"][2]]
        before_hdr = [c.value for c in load_workbook(p)["ZR1"][1]]

        # check_stock_excel 在函数体内部 from shein_scraper import _ensure_chrome，
        # 所以打桩模块属性即可拦住它。
        launched = []
        orig = shein_scraper._ensure_chrome

        def boom(*a, **k):
            launched.append(1)
            raise AssertionError("新模板不该启动 Chrome")

        shein_scraper._ensure_chrome = boom
        try:
            check_stock.check_stock_excel(p)
        finally:
            shein_scraper._ensure_chrome = orig

        assert launched == [], "Chrome 被启动了"
        after_row = [c.value for c in load_workbook(p)["ZR1"][2]]
        after_hdr = [c.value for c in load_workbook(p)["ZR1"][1]]
        assert before_row == after_row, f"数据行被改动\n{before_row}\n{after_row}"
        assert after_hdr == EXPECTED_HEADERS, after_hdr   # 没被塞 Stock/Last Checked
        assert before_hdr == after_hdr
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    test_new_template_is_untouched_and_chrome_not_launched()
    print("ALL PASS")
