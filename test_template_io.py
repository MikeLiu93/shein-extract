"""Test: run_excel.py 的中文模板读写层。
Run: python test_template_io.py
"""
from openpyxl import Workbook

import run_excel
from run_excel import (_sheet_matches_template, _read_pending_rows,
                       _write_result_row, EXPECTED_HEADERS,
                       COL_DATE, COL_STATUS, COL_WEB_PRICE, COL_SHEIN_TITLE,
                       COL_EBAY_TITLE, COL_EBAY_PRICE, COL_STOCK)


def _template_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "ZR1"
    for c, h in enumerate(EXPECTED_HEADERS, 1):
        ws.cell(1, c).value = h
    return wb, ws


def test_expected_headers_are_the_thirteen_chinese_names():
    assert EXPECTED_HEADERS == [
        "编号", "链接", "原价", "运费", "变体",
        "日期", "状态", "图片", "希音价格", "希音标题",
        "eBay标题", "eBay价格", "库存"], EXPECTED_HEADERS


def test_sheet_matches_only_when_b1_is_lianjie():
    wb, ws = _template_sheet()
    assert _sheet_matches_template(ws) is True
    ws.cell(1, 2).value = "Website"
    assert _sheet_matches_template(ws) is False


def test_old_english_sheet_is_rejected():
    wb = Workbook()
    ws = wb.active
    for c, h in enumerate(["Seq", "Website", "Price", "Date", "Status"], 1):
        ws.cell(1, c).value = h
    assert _sheet_matches_template(ws) is False
    assert _read_pending_rows(ws) == []


def test_reads_only_rows_with_url_and_empty_date_status():
    wb, ws = _template_sheet()
    ws.append([1, "http://a", 10.0, 2.99, "Black / M"])                  # 待处理
    ws.append([2, "http://b", None, None, None])                        # 待处理，C/D/E 空
    ws.append([3, "http://c", None, None, None, "2026-08-01", "Done"])   # 已完成
    ws.append([4, None, None, None, None])                              # 没链接
    ws.append([5, "http://e", None, None, None, "2026-08-01", None])     # 只有日期
    ws.append([6, "http://f", None, None, None, None, "Failed"])         # 只有状态
    pending = _read_pending_rows(ws)
    assert [p["seq"] for p in pending] == [1, 2], pending
    assert pending[0] == {"row": 2, "seq": 1, "url": "http://a",
                          "price": 10.0, "shipping": 2.99,
                          "variant_filter": "Black / M"}
    assert pending[1]["price"] is None
    assert pending[1]["shipping"] is None
    assert pending[1]["variant_filter"] == ""


def test_skips_rows_with_bad_seq_or_bad_numbers():
    wb, ws = _template_sheet()
    ws.append([None, "http://a", None, None, None])        # 没编号
    ws.append(["abc", "http://b", None, None, None])       # 编号非数字
    ws.append([3, "http://c", "十块", None, None])          # 原价非数字
    ws.append([4, "http://d", None, "免运费", None])         # 运费非数字
    ws.append([5, "http://e", 9.9, 0, None])               # 好行
    pending = _read_pending_rows(ws)
    assert [p["seq"] for p in pending] == [5], pending
    assert pending[0]["shipping"] == 0.0


def test_url_is_stripped():
    wb, ws = _template_sheet()
    ws.append([1, "  http://a  ", None, None, None])
    assert _read_pending_rows(ws)[0]["url"] == "http://a"


def test_write_result_row_writes_f_through_m():
    wb, ws = _template_sheet()
    ws.append([1, "http://a", None, None, None])
    _write_result_row(ws, row=2, date="2026-08-02", status="Done",
                      picture_path=None, web_price="9.99–14.99",
                      shein_title="Raw Title", ebay_title="eBay Title",
                      ebay_price=18.19, stock="Black-M: 12")
    assert ws.cell(2, COL_DATE).value == "2026-08-02"
    assert ws.cell(2, COL_STATUS).value == "Done"
    assert ws.cell(2, COL_WEB_PRICE).value == "9.99–14.99"
    assert ws.cell(2, COL_SHEIN_TITLE).value == "Raw Title"
    assert ws.cell(2, COL_EBAY_TITLE).value == "eBay Title"
    assert ws.cell(2, COL_EBAY_PRICE).value == 18.19
    assert ws.cell(2, COL_STOCK).value == "Black-M: 12"


def test_write_result_row_failed_writes_only_date_and_status():
    wb, ws = _template_sheet()
    ws.append([1, "http://a", None, None, None])
    _write_result_row(ws, row=2, date="2026-08-02", status="Failed")
    assert ws.cell(2, COL_DATE).value == "2026-08-02"
    assert ws.cell(2, COL_STATUS).value == "Failed"
    for col in (COL_WEB_PRICE, COL_SHEIN_TITLE, COL_EBAY_TITLE,
                COL_EBAY_PRICE, COL_STOCK):
        assert ws.cell(2, col).value is None, col


def test_numeric_web_price_stays_numeric():
    wb, ws = _template_sheet()
    ws.append([1, "http://a", None, None, None])
    _write_result_row(ws, row=2, date="2026-08-02", status="Done",
                      web_price=12.99)
    v = ws.cell(2, COL_WEB_PRICE).value
    assert isinstance(v, float) and v == 12.99, repr(v)


# ── process_excel 端到端（打桩 scrape_shein，不碰浏览器）────────────────────
import shutil
import tempfile
from pathlib import Path


def _fake_results(seqs):
    return [{
        "seq_num": s, "status": "OK", "sku": f"sku{s}",
        "title": f"Title {s}", "original_title": f"Raw Title {s}",
        "ebay_title": f"eBay Title {s}", "ebay_price": 18.19 + s,
        "web_price_display": "9.99–14.99", "stock_summary": f"Black-M: {s}",
        "first_image_path": "",
    } for s in seqs]


def _run_process_excel(rows, results_factory, tmp_root):
    """在临时目录里建模板 xlsx + 输出目录，打桩 scrape_shein 后跑 process_excel。
    返回 (加载回来的 ws, 传给 scrape_shein 的参数字典)。"""
    from openpyxl import load_workbook

    wb, ws = _template_sheet()
    for row in rows:
        ws.append(row)
    xlsx = Path(tmp_root) / "input.xlsx"
    wb.save(xlsx)

    captured = {}

    def fake_scrape(urls, output=None, seq_list=None, price_list=None,
                    shipping_list=None, variant_filter_list=None, **kw):
        captured.update(urls=urls, output=output, seq_list=seq_list,
                        price_list=price_list, shipping_list=shipping_list,
                        variant_filter_list=variant_filter_list)
        # 造出 seq 文件夹，process_excel 用它判断"有没有产出文件"
        for s in seq_list:
            d = Path.cwd() / str(s)
            d.mkdir(parents=True, exist_ok=True)
            (d / "img_001.webp").write_bytes(b"x")
        return results_factory(seq_list)

    orig_scrape = run_excel.scrape_shein
    orig_out = run_excel.OUTPUT_ROOT
    run_excel.scrape_shein = fake_scrape
    run_excel.OUTPUT_ROOT = Path(tmp_root) / "out"
    try:
        run_excel.process_excel(xlsx)
    finally:
        run_excel.scrape_shein = orig_scrape
        run_excel.OUTPUT_ROOT = orig_out

    return load_workbook(xlsx)["ZR1"], captured


def test_process_excel_fills_f_through_m_and_passes_cde():
    tmp = tempfile.mkdtemp()
    try:
        ws2, cap = _run_process_excel(
            [[1, "http://a", 10.0, 0.0, "Black / M"],
             [2, "http://b", None, None, None]],
            _fake_results, tmp)
        assert cap["urls"] == ["http://a", "http://b"]
        assert cap["seq_list"] == [1, 2]
        assert cap["price_list"] == [10.0, None]
        assert cap["shipping_list"] == [0.0, None]
        assert cap["variant_filter_list"] == ["Black / M", ""]
        assert ws2.cell(2, COL_STATUS).value == "Done"
        assert ws2.cell(2, COL_SHEIN_TITLE).value == "Raw Title 1"
        assert ws2.cell(2, COL_EBAY_TITLE).value == "eBay Title 1"
        assert ws2.cell(2, COL_EBAY_PRICE).value == 19.19
        assert ws2.cell(2, COL_WEB_PRICE).value == "9.99–14.99"
        assert ws2.cell(2, COL_STOCK).value == "Black-M: 1"
        assert ws2.cell(3, COL_STATUS).value == "Done"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_process_excel_marks_delisted_and_failed():
    tmp = tempfile.mkdtemp()
    try:
        def results(seqs):
            return [{"seq_num": 1, "status": "DELISTED"},
                    {"seq_num": 2, "status": "ERROR: boom"}]

        ws2, _ = _run_process_excel(
            [[1, "http://a", None, None, None],
             [2, "http://b", None, None, None]], results, tmp)
        assert ws2.cell(2, COL_STATUS).value == "Delisted"
        assert ws2.cell(3, COL_STATUS).value == "Failed"
        # 失败/下架的行只写日期和状态
        assert ws2.cell(2, COL_EBAY_PRICE).value is None
        assert ws2.cell(3, COL_EBAY_PRICE).value is None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_process_excel_skips_non_template_sheet():
    tmp = tempfile.mkdtemp()
    try:
        from openpyxl import load_workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "OldSheet"
        for c, h in enumerate(["Seq", "Website", "Price", "Date", "Status"], 1):
            ws.cell(1, c).value = h
        ws.append([1, "http://a", 9.9, None, None])
        xlsx = Path(tmp) / "old.xlsx"
        wb.save(xlsx)

        called = []
        orig = run_excel.scrape_shein
        run_excel.scrape_shein = lambda *a, **k: called.append(1) or []
        try:
            run_excel.process_excel(xlsx)
        finally:
            run_excel.scrape_shein = orig
        assert called == [], "旧英文表不该被抓取"
        ws2 = load_workbook(xlsx)["OldSheet"]
        assert ws2.cell(2, 4).value is None and ws2.cell(2, 5).value is None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ── 多 sheet：一次跑完，且每个 sheet 存一次不能把图片存废 ────────────────────

def _multi_sheet_xlsx(tmp_root, sheet_names, rows_per_sheet=2):
    """建一个多 sheet 的模板 xlsx，每个 sheet 都有待处理行。"""
    wb = Workbook()
    wb.remove(wb.active)
    for name in sheet_names:
        ws = wb.create_sheet(name)
        for c, h in enumerate(EXPECTED_HEADERS, 1):
            ws.cell(1, c).value = h
        for i in range(rows_per_sheet):
            ws.append([i + 1, f"http://{name}/{i}", None, None, None])
    xlsx = Path(tmp_root) / "多店铺.xlsx"
    wb.save(xlsx)
    return xlsx


def _run_multi(xlsx, tmp_root, results_for):
    """打桩 scrape_shein 跑 process_excel；每条结果都带一张真图片，
    这样每个 sheet 保存时都会往工作簿里塞 openpyxl Image。"""
    from PIL import Image as PILImage

    img = Path(tmp_root) / "pic.webp"
    PILImage.new("RGB", (80, 100), (10, 120, 200)).save(img, format="WEBP")

    def fake_scrape(urls, output=None, seq_list=None, **kw):
        for s in seq_list:
            d = Path.cwd() / str(s)
            d.mkdir(parents=True, exist_ok=True)
            (d / "img_001.webp").write_bytes(b"x")
        return results_for(seq_list, str(img))

    orig_scrape, orig_out = run_excel.scrape_shein, run_excel.OUTPUT_ROOT
    run_excel.scrape_shein = fake_scrape
    run_excel.OUTPUT_ROOT = Path(tmp_root) / "out"
    try:
        run_excel.process_excel(Path(xlsx))
    finally:
        run_excel.scrape_shein = orig_scrape
        run_excel.OUTPUT_ROOT = orig_out


def test_process_excel_runs_every_template_sheet():
    """三个 sheet 必须全跑完。

    以前第二个 sheet 保存时炸 ValueError: I/O operation on closed file
    （openpyxl 存一次就把图片的 BytesIO close 掉了），异常直接冲出 sheet
    循环，后面的 sheet 一个都没跑（2026-08-17 现场）。
    """
    from openpyxl import load_workbook
    tmp = tempfile.mkdtemp()
    try:
        names = ["Test11", "Test12", "Test13"]
        xlsx = _multi_sheet_xlsx(tmp, names)

        def results(seqs, img_path):
            out = _fake_results(seqs)
            for r in out:
                r["first_image_path"] = img_path
            return out

        _run_multi(xlsx, tmp, results)

        wb2 = load_workbook(xlsx)
        for name in names:
            ws = wb2[name]
            for row in (2, 3):
                assert ws.cell(row, COL_STATUS).value == "Done", \
                    f"{name} 第 {row} 行没跑到: {ws.cell(row, COL_STATUS).value!r}"
            assert len(ws._images) == 2, (name, len(ws._images))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_captcha_abort_stops_the_round_and_leaves_rows_pending():
    """人机验证等超时 → 已抓到的写回，没轮到的行保持空白，后面的 sheet 不跑。"""
    from openpyxl import load_workbook
    from shein_scraper import CaptchaAbort
    tmp = tempfile.mkdtemp()
    try:
        xlsx = _multi_sheet_xlsx(tmp, ["S1", "S2"])

        def results(seqs, img_path):
            done = _fake_results([seqs[0]])
            done[0]["first_image_path"] = img_path
            raise CaptchaAbort(
                "人机验证等待超过 10 分钟没人处理，本轮结束",
                records=done + [{"seq_num": seqs[1], "status": "CAPTCHA_ABORTED"}],
            )

        _run_multi(xlsx, tmp, results)

        wb2 = load_workbook(xlsx)
        s1 = wb2["S1"]
        assert s1.cell(2, COL_STATUS).value == "Done"       # 抓到的写回去了
        assert s1.cell(3, COL_DATE).value is None           # 没轮到的留空
        assert s1.cell(3, COL_STATUS).value is None
        s2 = wb2["S2"]                                      # 整个 sheet 没跑
        assert s2.cell(2, COL_STATUS).value is None
        assert s2.cell(3, COL_STATUS).value is None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ── 备份：绝不能挡住开跑 ────────────────────────────────────────────────────

def test_backup_falls_back_to_input_dir_when_backup_dir_unusable():
    tmp = tempfile.mkdtemp()
    try:
        xlsx = Path(tmp) / "总表.xlsx"
        wb, _ = _template_sheet()
        wb.save(xlsx)

        orig = run_excel.BACKUP_DIR
        run_excel.BACKUP_DIR = Path(r"Z:\不存在的备份盘\总表")
        try:
            run_excel.backup_excel(xlsx)      # 不该抛
        finally:
            run_excel.BACKUP_DIR = orig

        copies = [p.name for p in Path(tmp).glob("总表_*.xlsx")]
        assert len(copies) == 1, copies
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_backup_dir_defaults_to_input_dir():
    import config
    assert config.BACKUP_DIR == config.SUBMITTED_DIR or \
        "SHEIN_BACKUP_DIR" in __import__("os").environ or \
        "SHEIN_BACKUP_BASE" in __import__("os").environ, config.BACKUP_DIR


if __name__ == "__main__":
    test_expected_headers_are_the_thirteen_chinese_names()
    test_sheet_matches_only_when_b1_is_lianjie()
    test_old_english_sheet_is_rejected()
    test_reads_only_rows_with_url_and_empty_date_status()
    test_skips_rows_with_bad_seq_or_bad_numbers()
    test_url_is_stripped()
    test_write_result_row_writes_f_through_m()
    test_write_result_row_failed_writes_only_date_and_status()
    test_numeric_web_price_stays_numeric()
    test_process_excel_fills_f_through_m_and_passes_cde()
    test_process_excel_marks_delisted_and_failed()
    test_process_excel_skips_non_template_sheet()
    test_process_excel_runs_every_template_sheet()
    test_captcha_abort_stops_the_round_and_leaves_rows_pending()
    test_backup_falls_back_to_input_dir_when_backup_dir_unusable()
    test_backup_dir_defaults_to_input_dir()
    print("ALL PASS")
